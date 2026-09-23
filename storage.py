# it handle Excel and sync state management, including reading/writing spreadsheets, normalizing values, and tracking sync state for tickets. It defines constants for required sheet columns, editable columns, and read-only columns. It provides functions to ensure the workbook exists with the correct structure, read and write rows to the spreadsheet, update specific cells, and manage sync state in a JSON file. Additionally, it includes helpers for comparing values and building representations of ticket states for conflict detection.

from __future__ import annotations

import json
import os
import tempfile
from copy import deepcopy
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import Workbook, load_workbook


# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

PROJECT_DIR = Path(__file__).resolve().parent

DEFAULT_XLSX_PATH = PROJECT_DIR / "fieldserve.xlsx"
DEFAULT_STATE_PATH = PROJECT_DIR / "sync_state.json"


# ---------------------------------------------------------------------------
# Spreadsheet contract
# ---------------------------------------------------------------------------

# These are the exact required columns from the assignment.
SHEET_COLUMNS = [
    "Ticket ID",
    "Subject",
    "Status",
    "Engineer",
    "Manager",
    "Purchase Order",
    "Delivery Note",
    "Work Report",
    "Invoice",
    "Stage",
    "Action Owner",
]


# Fields that humans are allowed to edit through sheet.py.
EDITABLE_COLUMNS = {
    "Status",
    "Engineer",
    "Manager",
    "Purchase Order",
    "Delivery Note",
    "Work Report",
    "Invoice",
}


# API fields corresponding to editable sheet columns.
SHEET_TO_API_FIELD = {
    "Status": "status",
    "Engineer": "engineer_id",
    "Manager": "manager",
    "Purchase Order": "purchase_order",
    "Delivery Note": "delivery_note",
    "Work Report": "work_report",
    "Invoice": "invoice",
}


# Fields whose values are controlled by the API/workflow and should
# never be treated as ordinary human-editable fields.
READ_ONLY_COLUMNS = {
    "Ticket ID",
    "Subject",
    "Stage",
    "Action Owner",
}


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------

def normalize_cell_value(value: Any) -> Any:
    """
    Normalize values read from Excel.

    Empty strings and whitespace-only strings become None.
    Other values are preserved.
    """
    if value is None:
        return None

    if isinstance(value, str):
        stripped = value.strip()

        if stripped == "":
            return None

        return stripped

    return value


def normalize_ticket_id(value: Any) -> Optional[int]:
    """Convert a spreadsheet ticket ID to an integer."""
    if value is None or value == "":
        return None

    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def values_equal(left: Any, right: Any) -> bool:
    """
    Compare values in a way suitable for sync conflict detection.

    Excel may return numeric values differently from API JSON, so
    ticket IDs and status values are normalized when possible.
    """
    left = normalize_cell_value(left)
    right = normalize_cell_value(right)

    if left is None and right is None:
        return True

    if isinstance(left, bool) or isinstance(right, bool):
        return left == right

    # Numeric strings and numeric values should compare equally.
    try:
        if left is not None and right is not None:
            if float(left) == float(right):
                return True
    except (TypeError, ValueError):
        pass

    return left == right


# ---------------------------------------------------------------------------
# Excel workbook management
# ---------------------------------------------------------------------------

def ensure_workbook(
    xlsx_path: Path | str = DEFAULT_XLSX_PATH,
) -> Path:
    """
    Ensure that the workbook exists and has the required header row.
    """
    path = Path(xlsx_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if not path.exists():
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Tickets"

        for column_index, column_name in enumerate(
            SHEET_COLUMNS,
            start=1,
        ):
            worksheet.cell(
                row=1,
                column=column_index,
                value=column_name,
            )

        workbook.save(path)
        return path

    # Existing workbook: make sure the required sheet/header exists.
    workbook = load_workbook(path)
    worksheet = workbook["Tickets"] if "Tickets" in workbook.sheetnames else workbook.active

    existing_headers = [
        worksheet.cell(row=1, column=index).value
        for index in range(1, len(SHEET_COLUMNS) + 1)
    ]

    if existing_headers != SHEET_COLUMNS:
        # If the workbook is empty, initialize it.
        if all(value is None for value in existing_headers):
            for column_index, column_name in enumerate(
                SHEET_COLUMNS,
                start=1,
            ):
                worksheet.cell(
                    row=1,
                    column=column_index,
                    value=column_name,
                )
        else:
            raise ValueError(
                "fieldserve.xlsx does not have the required "
                "sheet column structure."
            )

        workbook.save(path)

    return path


def _header_map(worksheet: Any) -> Dict[str, int]:
    """Return column name -> 1-based Excel column number."""
    mapping: Dict[str, int] = {}

    for column_index, value in enumerate(
        worksheet[1],
        start=1,
    ):
        if value is None:
            continue

        mapping[str(value.value).strip()] = column_index

    return mapping

def read_sheet_rows(
    xlsx_path: Path | str = DEFAULT_XLSX_PATH,
) -> List[Dict[str, Any]]:
    """
    Read all ticket rows from the workbook.

    Returns one dictionary per row, keyed by the exact sheet column names.
    """
    path = ensure_workbook(xlsx_path)

    workbook = load_workbook(
        path,
        data_only=False,
    )

    worksheet = workbook["Tickets"]
    headers = _header_map(worksheet)

    missing_columns = [
        column
        for column in SHEET_COLUMNS
        if column not in headers
    ]

    if missing_columns:
        raise ValueError(
            "Workbook is missing required columns: "
            + ", ".join(missing_columns)
        )

    rows: List[Dict[str, Any]] = []

    for row_number in range(2, worksheet.max_row + 1):
        row: Dict[str, Any] = {}

        for column_name in SHEET_COLUMNS:
            column_index = headers[column_name]

            row[column_name] = normalize_cell_value(
                worksheet.cell(
                    row=row_number,
                    column=column_index,
                ).value
            )

        # Completely empty rows are ignored.
        if all(
            value is None
            for value in row.values()
        ):
            continue

        rows.append(row)

    return rows


def read_sheet_index(
    xlsx_path: Path | str = DEFAULT_XLSX_PATH,
) -> Dict[int, Dict[str, Any]]:
    """
    Read spreadsheet rows indexed by Ticket ID.
    """
    rows = read_sheet_rows(xlsx_path)

    result: Dict[int, Dict[str, Any]] = {}

    for row in rows:
        ticket_id = normalize_ticket_id(
            row.get("Ticket ID")
        )

        if ticket_id is None:
            continue

        result[ticket_id] = row

    return result


def write_sheet_rows(
    rows: Iterable[Dict[str, Any]],
    xlsx_path: Path | str = DEFAULT_XLSX_PATH,
) -> None:
    """
    Replace the ticket data in the workbook with the supplied rows.

    The header remains exactly the required assignment contract.
    """
    path = ensure_workbook(xlsx_path)

    # Load existing workbook so we can preserve workbook structure.
    workbook = load_workbook(path)

    if "Tickets" in workbook.sheetnames:
        worksheet = workbook["Tickets"]
    else:
        worksheet = workbook.active
        worksheet.title = "Tickets"

    # Ensure exact headers.
    for column_index, column_name in enumerate(
        SHEET_COLUMNS,
        start=1,
    ):
        worksheet.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

    # Remove existing data rows.
    if worksheet.max_row > 1:
        worksheet.delete_rows(
            2,
            worksheet.max_row - 1,
        )

    header_map = {
        column_name: index
        for index, column_name in enumerate(
            SHEET_COLUMNS,
            start=1,
        )
    }

    sorted_rows = sorted(
        rows,
        key=lambda row: (
            normalize_ticket_id(row.get("Ticket ID"))
            is None,
            normalize_ticket_id(row.get("Ticket ID"))
            or 0,
        ),
    )

    for row_number, row in enumerate(
        sorted_rows,
        start=2,
    ):
        for column_name in SHEET_COLUMNS:
            value = row.get(column_name)

            worksheet.cell(
                row=row_number,
                column=header_map[column_name],
                value=value,
            )

    # Basic usability formatting.
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Make long document URLs visible without changing their values.
    for column_name in (
        "Purchase Order",
        "Delivery Note",
        "Work Report",
        "Invoice",
    ):
        column_index = header_map[column_name]
        worksheet.column_dimensions[
            worksheet.cell(
                row=1,
                column=column_index,
            ).column_letter
        ].width = 45

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 55
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 25
    worksheet.column_dimensions["E"].width = 25
    worksheet.column_dimensions["J"].width = 22
    worksheet.column_dimensions["K"].width = 25

    # Atomic-ish save: write to a temporary file and replace the
    # existing workbook only after save succeeds.
    temporary_file: Optional[str] = None

    try:
        with tempfile.NamedTemporaryFile(
            suffix=".xlsx",
            delete=False,
            dir=str(path.parent),
        ) as temp:
            temporary_file = temp.name

        workbook.save(temporary_file)
        os.replace(temporary_file, path)

    finally:
        if temporary_file and os.path.exists(temporary_file):
            os.remove(temporary_file)


def update_sheet_cells(
    updates: Dict[int, Dict[str, Any]],
    xlsx_path: Path | str = DEFAULT_XLSX_PATH,
) -> None:
    """
    Update selected cells without rebuilding the entire workbook.

    `updates` format:

        {
            123456: {
                "Status": "Pending",
                "Manager": "Nefo Zagi",
            }
        }
    """
    path = ensure_workbook(xlsx_path)

    workbook = load_workbook(path)
    worksheet = workbook["Tickets"]

    headers = _header_map(worksheet)

    ticket_id_column = headers["Ticket ID"]

    ticket_rows: Dict[int, int] = {}

    for row_number in range(2, worksheet.max_row + 1):
        ticket_id = normalize_ticket_id(
            worksheet.cell(
                row=row_number,
                column=ticket_id_column,
            ).value
        )

        if ticket_id is not None:
            ticket_rows[ticket_id] = row_number

    for ticket_id, row_updates in updates.items():
        row_number = ticket_rows.get(ticket_id)

        if row_number is None:
            continue

        for column_name, value in row_updates.items():
            if column_name not in headers:
                raise ValueError(
                    f"Unknown sheet column: {column_name}"
                )

            worksheet.cell(
                row=row_number,
                column=headers[column_name],
                value=value,
            )

    workbook.save(path)


# ---------------------------------------------------------------------------
# Sync-state management
# ---------------------------------------------------------------------------

def empty_sync_state() -> Dict[str, Any]:
    """
    Return the initial sync-state structure.

    `tickets` stores the last synchronized API/sheet representation
    of editable fields.

    `server_updated_at` stores the latest API timestamp observed for
    each ticket.
    """
    return {
        "version": 1,
        "tickets": {},
        "server_updated_at": {},
    }


def load_sync_state(
    state_path: Path | str = DEFAULT_STATE_PATH,
) -> Dict[str, Any]:
    """Load sync_state.json, creating a valid empty state if necessary."""
    path = Path(state_path)

    if not path.exists():
        return empty_sync_state()

    try:
        with path.open(
            "r",
            encoding="utf-8",
        ) as file:
            data = json.load(file)
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(
            f"Unable to read sync state: {path}"
        ) from exc

    if not isinstance(data, dict):
        raise ValueError(
            "sync_state.json must contain a JSON object."
        )

    data.setdefault("version", 1)
    data.setdefault("tickets", {})
    data.setdefault("server_updated_at", {})

    return data


def save_sync_state(
    state: Dict[str, Any],
    state_path: Path | str = DEFAULT_STATE_PATH,
) -> None:
    """Atomically save sync_state.json."""
    path = Path(state_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    temporary_file: Optional[str] = None

    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            suffix=".json",
            delete=False,
            dir=str(path.parent),
        ) as temp:
            temporary_file = temp.name

            json.dump(
                state,
                temp,
                indent=2,
                ensure_ascii=False,
                sort_keys=True,
            )

            temp.write("\n")

        os.replace(
            temporary_file,
            path,
        )

    finally:
        if temporary_file and os.path.exists(temporary_file):
            os.remove(temporary_file)


def get_state_ticket(
    state: Dict[str, Any],
    ticket_id: int,
) -> Optional[Dict[str, Any]]:
    """Get the previous synchronized state for one ticket."""
    tickets = state.get("tickets", {})

    value = tickets.get(str(ticket_id))

    if not isinstance(value, dict):
        return None

    return deepcopy(value)


def set_state_ticket(
    state: Dict[str, Any],
    ticket_id: int,
    ticket_state: Dict[str, Any],
) -> None:
    """Store the latest synchronized state for one ticket."""
    state.setdefault("tickets", {})

    state["tickets"][str(ticket_id)] = deepcopy(
        ticket_state
    )


def delete_state_ticket(
    state: Dict[str, Any],
    ticket_id: int,
) -> None:
    """Remove one ticket from sync state."""
    tickets = state.setdefault(
        "tickets",
        {},
    )

    tickets.pop(
        str(ticket_id),
        None,
    )


# ---------------------------------------------------------------------------
# State representation
# ---------------------------------------------------------------------------

def build_sync_ticket_state(
    ticket: Dict[str, Any],
) -> Dict[str, Any]:
    """
    Build the portion of a ticket that matters for conflict detection.

    We intentionally store the API-side editable values rather than
    the complete raw ticket.
    """
    return {
        "status": ticket.get("status"),
        "engineer_id": ticket.get("engineer_id"),
        "manager": ticket.get("manager"),
        "purchase_order": ticket.get("purchase_order"),
        "delivery_note": ticket.get("delivery_note"),
        "work_report": ticket.get("work_report"),
        "invoice": ticket.get("invoice"),
        "subject": ticket.get("subject"),
        "updated_at": ticket.get("updated_at"),
    }


def build_sheet_ticket_state(
    row: Dict[str, Any],
) -> Dict[str, Any]:
    """
    Build the sheet representation relevant to conflict detection.
    """
    return {
        "status": row.get("Status"),
        "engineer": row.get("Engineer"),
        "manager": row.get("Manager"),
        "purchase_order": row.get("Purchase Order"),
        "delivery_note": row.get("Delivery Note"),
        "work_report": row.get("Work Report"),
        "invoice": row.get("Invoice"),
    }


# ---------------------------------------------------------------------------
# Conflict comparison helpers
# ---------------------------------------------------------------------------

def sheet_value_changed_since_state(
    current_row: Dict[str, Any],
    previous_state: Dict[str, Any],
    column_name: str,
) -> bool:
    """
    Determine whether a sheet cell changed since the last sync.
    """
    state_field = {
        "Status": "status",
        "Engineer": "engineer",
        "Manager": "manager",
        "Purchase Order": "purchase_order",
        "Delivery Note": "delivery_note",
        "Work Report": "work_report",
        "Invoice": "invoice",
    }.get(column_name)

    if state_field is None:
        return False

    current_value = current_row.get(column_name)
    previous_value = previous_state.get(
        state_field
    )

    return not values_equal(
        current_value,
        previous_value,
    )


def api_value_changed_since_state(
    current_ticket: Dict[str, Any],
    previous_state: Dict[str, Any],
    column_name: str,
) -> bool:
    """
    Determine whether the server value changed since the last sync.
    """
    api_field = SHEET_TO_API_FIELD.get(column_name)

    if api_field is None:
        return False

    current_value = current_ticket.get(api_field)
    previous_value = previous_state.get(
        api_field
    )

    return not values_equal(
        current_value,
        previous_value,
    )


# ---------------------------------------------------------------------------
# Export helpers
# ---------------------------------------------------------------------------

def rows_for_export(
    xlsx_path: Path | str = DEFAULT_XLSX_PATH,
) -> List[Dict[str, Any]]:
    """
    Return rows in the format required by `sheet.py export`.
    """
    return read_sheet_rows(xlsx_path)